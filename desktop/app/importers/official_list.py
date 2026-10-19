import csv
from datetime import date, datetime
from openpyxl import load_workbook

COLUMNS = ["card_id", "first_name", "last_name", "birth_date", "gender", "university", "faculty"]


def read_official_file(path: str) -> list[dict]:
    if path.lower().endswith(".csv"):
        with open(path, newline="", encoding="utf-8-sig") as fh:
            raw_rows = list(csv.DictReader(fh))
    else:
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb.active
            first_row = next(ws.iter_rows(min_row=1, max_row=1), None) if ws is not None else None
            if not first_row:
                header = []
            else:
                header = [str(c.value).strip() if c.value is not None else "" for c in first_row]
            raw_rows = []
            if header:
                for row in ws.iter_rows(min_row=2, values_only=True):
                    raw_rows.append({header[i]: row[i] for i in range(min(len(header), len(row)))})
        finally:
            wb.close()

    cleaned: list[dict] = []
    for raw in raw_rows:
        r = {str(k).strip().lower().replace(" ", "_"): v for k, v in raw.items() if k}
        card = str(r.get("card_id") or "").strip()
        if not card:
            continue

        b = r.get("birth_date")
        if isinstance(b, (date, datetime)):
            birth_date = b.strftime("%Y-%m-%d")
        else:
            b_str = str(b or "").strip()
            if b_str:
                parsed_date = None
                for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
                    try:
                        parsed_date = datetime.strptime(b_str, fmt).strftime("%Y-%m-%d")
                        break
                    except ValueError:
                        pass
                birth_date = parsed_date if parsed_date else b_str
            else:
                birth_date = None

        gender_raw = str(r.get("gender") or "").strip().lower()
        if gender_raw in ("m", "homme", "masculin", "male"):
            gender_val = "male"
        elif gender_raw in ("f", "femme", "feminin", "féminin", "female"):
            gender_val = "female"
        else:
            gender_val = gender_raw or None

        cleaned.append({
            "card_id": card,
            "first_name": str(r.get("first_name") or "").strip(),
            "last_name": str(r.get("last_name") or "").strip(),
            "birth_date": birth_date,
            "gender": gender_val,
            "university": str(r.get("university") or "").strip() or None,
            "faculty": str(r.get("faculty") or "").strip() or None,
        })
    return cleaned

